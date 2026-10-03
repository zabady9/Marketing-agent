"""Turns an uploaded ChatAttachment into something the chat agent can read.
Each file goes down one of three paths (ChatAttachment.kind):

- "gemini_file": images, PDFs, audio and video go to the Gemini File API and
  are replayed to the model as native multimodal parts. These uploads expire
  after 48h, so ensure_gemini_file re-uploads them when an older session's
  history is replayed.
- "text": office documents (DOCX/XLSX/PPTX) and anything that decodes as
  text (CSV, JSON, markdown, source code, logs...) are extracted to plain
  text locally. Cheaper than a Gemini upload, and it never expires.
- "metadata_only": everything else (archives, executables, unknown binaries).
  The agent is told only the file's name, type and size.

A failure on any path marks the attachment "failed" with an error message.
It can still be sent; the agent just sees it as metadata only.
"""

from __future__ import annotations

import asyncio
import csv
import io
import logging
import mimetypes
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import BinaryIO

import filetype
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import ChatAttachment
from app.services import upload_storage

logger = logging.getLogger(__name__)

_GEMINI_NATIVE_PREFIXES = ("image/", "audio/", "video/")
_GEMINI_NATIVE_TYPES = {"application/pdf"}

_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
_OFFICE_BY_EXT = {".docx": _DOCX, ".xlsx": _XLSX, ".xlsm": _XLSX, ".pptx": _PPTX}

# How much of the head of a file is read to sniff its type / text-ness.
_SNIFF_BYTES = 8192
_XLSX_MAX_ROWS_PER_SHEET = 2000
# Gemini needs a few seconds to tens of seconds to make a video ACTIVE.
_GEMINI_ACTIVE_TIMEOUT_SECONDS = 600
_GEMINI_POLL_SECONDS = 3
# Re-upload a Gemini file this long before it actually expires, so it can't
# expire in the middle of a turn.
_GEMINI_EXPIRY_MARGIN = timedelta(hours=1)


def resolve_content_type(filename: str, declared: str, head: bytes) -> str:
    """The best-guess MIME type: the extension wins for OOXML (filetype
    reports those as plain zip), then magic bytes, then the browser's
    declared type, then the extension via mimetypes."""
    ext = Path(filename).suffix.lower()
    if ext in _OFFICE_BY_EXT:
        return _OFFICE_BY_EXT[ext]
    kind = filetype.guess(head) if head else None
    if kind is not None:
        return kind.mime
    if declared and declared != "application/octet-stream":
        return declared
    guessed, _ = mimetypes.guess_type(filename)
    return guessed or "application/octet-stream"


def _is_gemini_native(content_type: str) -> bool:
    if content_type == "image/svg+xml":  # XML text, which Gemini won't take as an image
        return False
    return content_type in _GEMINI_NATIVE_TYPES or content_type.startswith(_GEMINI_NATIVE_PREFIXES)


def _looks_like_text(head: bytes) -> bool:
    if not head or b"\x00" in head:
        return False
    try:
        head.decode("utf-8")
        return True
    except UnicodeDecodeError as exc:
        # The sniff window may have cut a multi-byte character in half.
        return exc.start >= len(head) - 4


def _truncate(text: str, limit: int) -> str:
    if len(text) <= limit:
        return text
    return text[:limit] + f"\n\n[Truncated — showing the first {limit:,} of {len(text):,} characters.]"


def _extract_docx(stream: BinaryIO) -> str:
    import docx

    document = docx.Document(stream)
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _extract_xlsx(stream: BinaryIO) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(stream, read_only=True, data_only=True)
    out = io.StringIO()
    writer = csv.writer(out)
    for sheet in workbook.worksheets:
        out.write(f"## Sheet: {sheet.title}\n")
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= _XLSX_MAX_ROWS_PER_SHEET:
                out.write(f"[... rows beyond {_XLSX_MAX_ROWS_PER_SHEET} omitted]\n")
                break
            if any(v is not None for v in row):
                writer.writerow(["" if v is None else v for v in row])
        out.write("\n")
    workbook.close()
    return out.getvalue()


def _extract_pptx(stream: BinaryIO) -> str:
    from pptx import Presentation

    presentation = Presentation(stream)
    parts = []
    for n, slide in enumerate(presentation.slides, start=1):
        parts.append(f"## Slide {n}")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append(f"Notes: {slide.notes_slide.notes_text_frame.text}")
    return "\n".join(parts)


def _extract_text(content_type: str, storage_path: str, max_chars: int) -> str:
    with upload_storage.open_read(storage_path) as stream:
        if content_type == _DOCX:
            text = _extract_docx(stream)
        elif content_type == _XLSX:
            text = _extract_xlsx(stream)
        elif content_type == _PPTX:
            text = _extract_pptx(stream)
        else:
            # Read a little past the limit so _truncate can still say the
            # file was cut, without pulling a multi-GB log into memory.
            text = stream.read(max_chars * 4 + 4).decode("utf-8", errors="replace")
    return _truncate(text, max_chars)


def _upload_to_gemini(storage_path: str, content_type: str, display_name: str) -> tuple[str, datetime | None]:
    """Blocking: a resumable upload to the Gemini File API, streamed from
    storage, then polls until the file is ACTIVE. Returns (uri, expires_at)."""
    from google import genai

    client = genai.Client(api_key=get_settings().google_api_key)
    with upload_storage.open_read(storage_path) as stream:
        uploaded = client.files.upload(
            file=stream, config={"mime_type": content_type, "display_name": display_name[:128]}
        )
    deadline = time.monotonic() + _GEMINI_ACTIVE_TIMEOUT_SECONDS
    while uploaded.state is not None and uploaded.state.name == "PROCESSING":
        if time.monotonic() > deadline:
            raise TimeoutError("Gemini is still processing the file after 10 minutes.")
        time.sleep(_GEMINI_POLL_SECONDS)
        uploaded = client.files.get(name=uploaded.name)
    if uploaded.state is not None and uploaded.state.name == "FAILED":
        raise RuntimeError(f"Gemini could not process the file: {uploaded.error}")
    expires_at = uploaded.expiration_time
    if expires_at is not None and expires_at.tzinfo is not None:
        expires_at = expires_at.astimezone(timezone.utc).replace(tzinfo=None)
    return uploaded.uri, expires_at


def _read_head(storage_path: str) -> bytes:
    with upload_storage.open_read(storage_path) as stream:
        return stream.read(_SNIFF_BYTES)


async def process_attachment(db: Session, attachment: ChatAttachment) -> None:
    """Classifies and prepares one uploaded attachment, then commits
    status="ready" or "failed". Blocking I/O (GCS reads, parsing, Gemini
    upload) runs on a worker thread."""
    settings = get_settings()
    try:
        head = await asyncio.to_thread(_read_head, attachment.storage_path)
        content_type = resolve_content_type(attachment.filename, attachment.content_type, head)
        attachment.content_type = content_type

        if _is_gemini_native(content_type):
            uri, expires_at = await asyncio.to_thread(
                _upload_to_gemini, attachment.storage_path, content_type, attachment.filename
            )
            attachment.kind = "gemini_file"
            attachment.gemini_file_uri = uri
            attachment.gemini_file_expires_at = expires_at
        elif content_type in (_DOCX, _XLSX, _PPTX) or _looks_like_text(head):
            attachment.extracted_text = await asyncio.to_thread(
                _extract_text, content_type, attachment.storage_path, settings.max_extracted_chars
            )
            attachment.kind = "text"
        else:
            attachment.kind = "metadata_only"
        attachment.status = "ready"
        attachment.error = None
    except Exception as exc:
        logger.exception("Processing attachment %s failed", attachment.id)
        attachment.kind = "metadata_only"
        attachment.status = "failed"
        attachment.error = str(exc)[:1000] or exc.__class__.__name__
    db.commit()


def _needs_gemini_refresh(attachment: ChatAttachment) -> bool:
    if attachment.kind != "gemini_file":
        return False
    if attachment.gemini_file_uri is None:
        return True
    expires_at = attachment.gemini_file_expires_at
    return expires_at is not None and expires_at - _GEMINI_EXPIRY_MARGIN <= datetime.utcnow()


async def ensure_gemini_files(db: Session, attachments: list[ChatAttachment]) -> None:
    """Re-uploads any expired (or nearly expired) Gemini file among these
    attachments from storage before their history is replayed. One that
    can't be re-uploaded falls back to metadata-only instead of breaking
    the turn."""
    stale = [a for a in attachments if _needs_gemini_refresh(a)]
    for attachment in stale:
        try:
            uri, expires_at = await asyncio.to_thread(
                _upload_to_gemini, attachment.storage_path, attachment.content_type, attachment.filename
            )
            attachment.gemini_file_uri = uri
            attachment.gemini_file_expires_at = expires_at
        except Exception as exc:
            logger.exception("Re-uploading attachment %s to Gemini failed", attachment.id)
            attachment.kind = "metadata_only"
            attachment.gemini_file_uri = None
            attachment.error = str(exc)[:1000] or exc.__class__.__name__
    if stale:
        db.commit()
